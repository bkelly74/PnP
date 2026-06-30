import openpyxl, warnings, json, csv, datetime, re
from collections import defaultdict, Counter
warnings.filterwarnings("ignore")
F="/root/.claude/uploads/f4169ad0-d1da-5751-9c56-85d42e4189cd/9d102f1d-N2C_Expanded_WBS.xlsx"
wb=openpyxl.load_workbook(F, data_only=True)
MULT={'Fortnightly':26,'Monthly':12,'Half yearly':2,'Daily(Mon-Fri)':253,'Quarterly':4,
'Twice a week':104,'Three times a week':156,'Weekly':52,'Yearly':1,'Daily(24x7)':365,'One off':1}
FTE_DAYS=220
DATA=['DLZ Years 1-5','Data Feeds Years 1-5','DARKSTAR OCD Years 1-2','AIP on ACHE Build Years 2-4',
'T&R on ACHE Site 2 Years 2-5','AIP on ACHE Site 2 Years 2-5','AIP on ACHE Site 3 Years 3-5',
'AIP on ACHE Site 1 Years 4-5','Local Site Years 1-2','STRATA Years 1-5','NEXUS Years 1-5',
'BLACK KNIGHT Years 1-5','Service Management Years 1-5','NOC Years 1-5','Service Desk Years 1-5',
'Security Years 1-5','Governance Years 1-5']
def isnum(v): return isinstance(v,(int,float)) and not isinstance(v,bool)
def num(v): return v if isnum(v) else None
def col(ws,r,c): return ws.cell(row=r,column=c).value
def parsed(v):
    if isinstance(v,datetime.datetime): return v
    if isinstance(v,str):
        m=re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})',v.strip())
        if m:
            d,mo,y=map(int,m.groups())
            try: return datetime.datetime(y,mo,d)
            except: return None
    return None
roles_ws=wb['Roles & SFIA']; master={}
for r in range(2,roles_ws.max_row+1):
    nm=roles_ws.cell(row=r,column=1).value
    if nm: master[str(nm).strip().lower()]=str(nm)
ISS=[]
def add(sheet,row,cat,sev,detail,ev=""): ISS.append(dict(severity=sev,category=cat,sheet=sheet,row=row,detail=detail,evidence=ev))

leaf=0; totals={'days':0.0,'tfte':0.0}; team_tot=Counter(); year_tot=Counter()
role_fte=Counter(); role_sfia=defaultdict(set); date_type=defaultdict(lambda:Counter())
ftemis=Counter()
for sh in DATA:
    ws=wb[sh]; blkcy=None; blkname=None
    for r in range(2,ws.max_row+1):
        nm=col(ws,r,1); fq=col(ws,r,2)
        if nm and isinstance(nm,str) and 'CONTRACT YEAR' in nm.upper():
            m=re.search(r'CONTRACT YEAR\s*(\d+)',nm.upper()); blkcy=int(m.group(1)) if m else None; blkname=nm.strip()
        if fq is None or (isinstance(fq,str) and not fq.strip()): continue
        leaf+=1
        O,ML,P,est,days,pfte,tfte=[num(col(ws,r,c)) for c in (3,4,5,6,7,8,9)]
        fkey=str(fq).strip(); mult=MULT.get(fkey)
        if mult is None: add(sh,r,'Unknown frequency','High',f"Frequency '{fq}' not in Mapping table",str(nm))
        if None not in (O,ML,P) and not (O<=ML<=P):
            add(sh,r,'3-point estimate ordering','High',f"O={O}, ML={ML}, P={P} not non-decreasing",str(nm))
        if None not in (O,ML,P,est):
            pert=(O+4*ML+P)/6
            if abs(pert-est)>max(1e-6,abs(pert)*1e-4): add(sh,r,'PERT estimate mismatch','Medium',f"Estimate={est} vs PERT={round(pert,4)}",str(nm))
        if None not in (est,mult,days) and abs(est*mult-days)>max(1e-6,abs(est*mult)*1e-4):
            add(sh,r,'Annual Effort (Days) mismatch','High',f"Days={round(days,3)} vs Est×{fkey}({mult})={round(est*mult,3)}",str(nm))
        if None not in (days,pfte) and abs(days/FTE_DAYS-pfte)>max(1e-6,abs(days/FTE_DAYS)*1e-4):
            add(sh,r,'%FTE mismatch','Medium',f"%FTE={round(pfte,4)} vs Days/220={round(days/FTE_DAYS,4)}",str(nm))
        roles=[]
        for rc,sc in ((10,11),(12,13),(14,15),(16,17),(18,19)):
            rv=col(ws,r,rc); sv=col(ws,r,sc)
            if rv and str(rv).strip():
                roles.append((str(rv).strip(), str(sv).strip() if sv else None)); role_sfia[str(rv).strip()].add(str(sv).strip() if sv else 'None')
        n=len(roles)
        if n==0: add(sh,r,'No role assigned','Medium',"Leaf task has effort but no role/resource",str(nm))
        if pfte not in (None,0) and tfte is not None:
            ratio=tfte/pfte; impl=round(ratio)
            if abs(ratio-impl)<0.02 and impl!=n and n>0:
                add(sh,r,'Total FTE vs role-count mismatch','Medium',f"Total FTE implies {impl} resources but {n} role(s) named",str(nm)); ftemis[sh]+=1
        for rv,sv in roles:
            if rv.lower() not in master: add(sh,r,'Role not in master list','Low',f"Role '{rv}' absent from 'Roles & SFIA'",str(nm))
            if not sv or sv=='None': add(sh,r,'Missing SFIA level','Low',f"Role '{rv}' has no SFIA level",str(nm))
        sd=col(ws,r,21); ed=col(ws,r,22)
        date_type[sh]['start_text']+= isinstance(sd,str); date_type[sh]['start_date']+= isinstance(sd,datetime.datetime)
        date_type[sh]['end_text']+= isinstance(ed,str); date_type[sh]['end_date']+= isinstance(ed,datetime.datetime)
        if isinstance(ed,str): add(sh,r,'Date stored as text','Medium',f"End Date is text '{ed}'",str(nm))
        if isinstance(sd,str): add(sh,r,'Date stored as text','Medium',f"Start Date is text '{sd}'",str(nm))
        psd,ped=parsed(sd),parsed(ed)
        if psd and ped and ped<psd: add(sh,r,'End before Start','High',f"End {ped.date()} < Start {psd.date()}",str(nm))
        if sd is None and ed is None: add(sh,r,'Missing dates','Low',"No Start or End date",str(nm))
        cy=col(ws,r,23)
        if cy is None: add(sh,r,'Missing Contract Year','Low',"Leaf row has no Contract Year",str(nm))
        elif isnum(cy) and blkcy is not None and cy!=blkcy:
            add(sh,r,'Contract Year vs block','Medium',f"Row Contract Year={cy} but block is '{blkname}' (Year {blkcy})",str(nm))
        if isnum(cy) and isinstance(col(ws,r,24),type(None)): add(sh,r,'Missing Team','Low',"Leaf row has no Team",str(nm))
        if tfte is not None and n==1 and tfte>1.0:
            add(sh,r,'Single-resource over 1 FTE','Medium',f"Needs {round(tfte,2)} FTE but 1 resource ({roles[0][0]}) named",str(nm))
        if days: totals['days']+=days; team_tot[col(ws,r,24) or sh]+=days; year_tot[cy if cy is not None else '?']+=days
        if tfte:
            totals['tfte']+=tfte
            for rv,sv in roles: role_fte[rv]+= tfte/max(n,1)
# exact dup
for sh in DATA:
    ws=wb[sh]; seen=defaultdict(list)
    for r in range(2,ws.max_row+1):
        fq=col(ws,r,2)
        if fq is None or (isinstance(fq,str) and not fq.strip()): continue
        key=tuple(str(col(ws,r,c)) for c in range(1,25)); seen[key].append(r)
    for k,rows in seen.items():
        if len(rows)>1: add(sh,rows[1],'Exact duplicate row','Low',f"Identical to row {rows[0]} (task '{k[0]}')",f"rows {rows}")
# SFIA inconsistency
for rv,levels in role_sfia.items():
    real={l for l in levels if l!='None'}
    if len(real)>1: add('Roles & SFIA',0,'SFIA level inconsistency','Low',f"Role '{rv}' appears at multiple SFIA levels: {sorted(real)}","")
# rollup not reflected
for rv,fte in role_fte.items():
    k=rv.lower()
    if k in master and fte>0.05:
        rr=next((r for r in range(2,roles_ws.max_row+1) if str(roles_ws.cell(row=r,column=1).value).strip().lower()==k),None)
        if rr:
            vals=[roles_ws.cell(row=rr,column=c).value for c in range(2,9)]
            if sum(v for v in vals if isnum(v))==0:
                add('Roles & SFIA',rr,'Rollup not reflected','Medium',f"Role '{rv}' ~{round(fte,2)} FTE in WBS but 0 in summary tab","")
sev={'High':0,'Medium':1,'Low':2}
ISS.sort(key=lambda x:(sev[x['severity']],x['category'],x['sheet'],x['row']))
with open('issues_final.csv','w',newline='') as fh:
    w=csv.DictWriter(fh,fieldnames=['severity','category','sheet','row','detail','evidence']); w.writeheader()
    for i in ISS: w.writerow(i)
json.dump(ISS,open('issues_final.json','w'),indent=1,default=str)
print("LEAF:",leaf,"ISSUES:",len(ISS))
print("\nBY SEVERITY:", {s:sum(1 for i in ISS if i['severity']==s) for s in ('High','Medium','Low')})
print("\nBY CATEGORY:")
for (s,c),v in sorted(Counter((i['severity'],i['category']) for i in ISS).items(),key=lambda x:(sev[x[0][0]],-x[1])):
    print(f"  [{s}] {c}: {v}")
print("\nDATE TYPE per sheet (start_text/start_date | end_text/end_date):")
for sh in DATA:
    d=date_type[sh]; print(f"  {sh}: S {d['start_text']}t/{d['start_date']}d | E {d['end_text']}t/{d['end_date']}d")
print("\nTotal FTE-mismatch rows per sheet:", dict(ftemis))
print("\nTOTAL annual days:",round(totals['days'],1)," TOTAL FTE:",round(totals['tfte'],2))
print("By year (days):", {k:round(v,1) for k,v in sorted(year_tot.items(),key=lambda x:str(x))})
print("\nTop roles by FTE:")
for rv,f in role_fte.most_common(12): print(f"  {rv}: {round(f,2)}")
